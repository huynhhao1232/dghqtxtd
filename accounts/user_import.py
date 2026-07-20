"""Import tài khoản viên chức từ Excel / CSV."""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass, field

from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from openpyxl import Workbook, load_workbook

from .models import User

# Cột mẫu (tiếng Việt)
TEMPLATE_HEADERS = [
    'Họ và tên',
    'Username',
    'Mật khẩu',
    'Vai trò',
    'Chức vụ',
]

HEADER_ALIASES = {
    'full_name': {
        'ho va ten', 'họ và tên', 'hoten', 'ho ten', 'full name', 'fullname', 'ten',
    },
    'username': {'username', 'ten dang nhap', 'tên đăng nhập', 'user'},
    'password': {'password', 'mat khau', 'mật khẩu', 'mk'},
    'role': {'role', 'vai tro', 'vai trò', 'quyen', 'quyền'},
    'position': {'position', 'chuc vu', 'chức vụ'},
}

ROLE_ALIASES = {
    User.ROLE_STAFF: {
        'staff', 'gv', 'giao vien', 'giáo viên', 'nhan vien', 'nhân viên',
        'giao vien / nhan vien', 'giáo viên / nhân viên',
    },
    User.ROLE_DEPARTMENT: {
        'department', 'dept', 'to', 'tổ', 'to chuyen mon', 'tổ chuyên môn',
        'truong to', 'trưởng tổ',
    },
    User.ROLE_DIRECTOR: {
        'director', 'manager', 'bgd', 'ban giam doc', 'ban giám đốc',
        'lanh dao', 'lãnh đạo', 'giam doc', 'giám đốc',
    },
}


@dataclass
class ImportRowResult:
    row_num: int
    username: str
    ok: bool
    message: str


@dataclass
class ImportResult:
    created: int = 0
    skipped: int = 0
    rows: list[ImportRowResult] = field(default_factory=list)

    @property
    def errors(self):
        return [r for r in self.rows if not r.ok]


def _normalize_header(value) -> str:
    text = unicodedata.normalize('NFC', str(value or '').strip().lower())
    text = text.replace('đ', 'd')
    text = ''.join(
        c for c in unicodedata.normalize('NFD', text)
        if not unicodedata.combining(c)
    )
    return text


def _map_headers(raw_headers: list) -> dict[str, int]:
    mapping = {}
    for idx, header in enumerate(raw_headers):
        key = _normalize_header(header)
        if not key:
            continue
        for field_name, aliases in HEADER_ALIASES.items():
            if key in aliases or key.replace(' ', '') in aliases:
                mapping[field_name] = idx
                break
    return mapping


def parse_role(value) -> str | None:
    raw = _normalize_header(value)
    if not raw:
        return None
    compact = raw.replace(' ', '')
    for role, aliases in ROLE_ALIASES.items():
        if raw in aliases or compact in {a.replace(' ', '') for a in aliases}:
            return role
    if raw in (User.ROLE_STAFF, User.ROLE_DEPARTMENT, User.ROLE_DIRECTOR):
        return raw
    return None


def build_import_template_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tai khoan'
    ws.append(TEMPLATE_HEADERS)
    ws.append([
        'Nguyễn Văn An',
        'nguyenvana',
        'MatKhau@123',
        'Giáo viên / Nhân viên',
        'Giáo viên',
    ])
    ws.append([
        'Trần Thị Bình',
        'tranthib',
        'MatKhau@123',
        'Tổ chuyên môn',
        'Trưởng tổ Toán',
    ])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
    return wb


def _cell_str(value) -> str:
    if value is None:
        return ''
    return str(value).strip()


def _rows_from_xlsx(file_obj) -> list[dict]:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []
    col_map = _map_headers(list(header_row))
    if 'username' not in col_map or 'password' not in col_map:
        raise ValueError(
            'File Excel thiếu cột bắt buộc. Cần ít nhất: Username, Mật khẩu.'
        )
    rows = []
    for line_no, cells in enumerate(rows_iter, start=2):
        if not cells or all(_cell_str(c) == '' for c in cells):
            continue
        row = {'_line': line_no}
        for field_name, idx in col_map.items():
            row[field_name] = _cell_str(cells[idx]) if idx < len(cells) else ''
        rows.append(row)
    return rows


def _rows_from_csv(file_obj) -> list[dict]:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        for encoding in ('utf-8-sig', 'utf-8', 'cp1258', 'latin-1'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('Không đọc được file CSV (mã hóa không hỗ trợ).')
    else:
        text = raw
    reader = csv.reader(io.StringIO(text))
    try:
        header_row = next(reader)
    except StopIteration:
        return []
    col_map = _map_headers(header_row)
    if 'username' not in col_map or 'password' not in col_map:
        raise ValueError(
            'File CSV thiếu cột bắt buộc. Cần ít nhất: Username, Mật khẩu.'
        )
    rows = []
    for line_no, cells in enumerate(reader, start=2):
        if not cells or all(_cell_str(c) == '' for c in cells):
            continue
        row = {'_line': line_no}
        for field_name, idx in col_map.items():
            row[field_name] = _cell_str(cells[idx]) if idx < len(cells) else ''
        rows.append(row)
    return rows


def read_import_file(uploaded_file) -> list[dict]:
    name = (uploaded_file.name or '').lower()
    if name.endswith('.csv'):
        return _rows_from_csv(uploaded_file)
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        return _rows_from_xlsx(uploaded_file)
    raise ValueError('Chỉ hỗ trợ file .xlsx hoặc .csv.')


def import_users_from_rows(rows: list[dict]) -> ImportResult:
    result = ImportResult()
    seen_usernames: set[str] = set()

    for row in rows:
        line = int(row.get('_line') or 0)
        username = (row.get('username') or '').strip()
        password = row.get('password') or ''
        full_name = (row.get('full_name') or '').strip()
        position = (row.get('position') or '').strip()
        role_raw = row.get('role') or ''

        if not username:
            result.skipped += 1
            result.rows.append(ImportRowResult(line, '—', False, 'Thiếu username.'))
            continue

        uname_key = username.lower()
        if uname_key in seen_usernames:
            result.skipped += 1
            result.rows.append(
                ImportRowResult(line, username, False, 'Username trùng trong file.')
            )
            continue
        seen_usernames.add(uname_key)

        if User.objects.filter(username__iexact=username).exists():
            result.skipped += 1
            result.rows.append(
                ImportRowResult(line, username, False, 'Username đã tồn tại trong hệ thống.')
            )
            continue

        if not password:
            result.skipped += 1
            result.rows.append(ImportRowResult(line, username, False, 'Thiếu mật khẩu.'))
            continue

        if not full_name:
            result.skipped += 1
            result.rows.append(ImportRowResult(line, username, False, 'Thiếu họ và tên.'))
            continue

        role = parse_role(role_raw)
        if not role:
            result.skipped += 1
            result.rows.append(
                ImportRowResult(
                    line,
                    username,
                    False,
                    f'Vai trò không hợp lệ: "{role_raw}".',
                ),
            )
            continue

        user = User(
            username=username,
            first_name=full_name,
            last_name='',
            position=position,
            role=role,
            is_manager=role == User.ROLE_DIRECTOR,
            is_staff=True,
            is_active=True,
            account_status=User.ACCOUNT_ACTIVE,
        )
        try:
            validate_password(password, user=user)
        except ValidationError as exc:
            result.skipped += 1
            result.rows.append(
                ImportRowResult(line, username, False, '; '.join(exc.messages)),
            )
            continue

        user.set_password(password)
        user.save()
        result.created += 1
        result.rows.append(
            ImportRowResult(line, username, True, 'Đã tạo tài khoản.'),
        )

    return result
